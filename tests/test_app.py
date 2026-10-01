from openpyxl import load_workbook

from .conftest import csrf, login


def test_login_required_and_login(client):
    assert client.get("/").status_code == 302
    response = login(client)
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/")
    assert client.get("/").status_code == 200


def test_invalid_login(client):
    response = login(client, password="wrong")
    assert b"Invalid username or password" in response.data


def test_regular_user_cannot_manage_accounts(client):
    login(client, "worker", "worker-password-123")
    assert client.get("/admin/users").status_code == 302


def test_admin_can_create_user_and_password_is_hashed(client, app):
    login(client)
    response = client.post("/admin/users", data={"csrf_token": csrf(client), "username": "new.user", "password": "secure-password-123", "role": "User"})
    assert response.status_code == 302
    wb = load_workbook(app.config["WORKBOOK_PATH"])
    row = next(row for row in wb["Users"].iter_rows(min_row=2, values_only=True) if row[0] == "new.user")
    assert row[1] != "secure-password-123"
    assert row[1].startswith("scrypt:")


def test_product_creation_and_barcode_lookup(client):
    login(client, "worker", "worker-password-123")
    response = client.post("/api/products", json={"barcode": "0123456789012", "name": "Widget", "category": "Parts", "quantity": 5, "location": "A-1"}, headers={"X-CSRF-Token": csrf(client)})
    assert response.status_code == 201
    assert response.json["product"]["product_id"] == "PRD-000001"
    lookup = client.get("/api/products/0123456789012")
    assert lookup.status_code == 200
    assert lookup.json["product"]["name"] == "Widget"


def test_quantity_updates_and_cannot_go_negative(client, app):
    login(client)
    token = csrf(client)
    client.post("/api/products", json={"barcode": "ABC123", "name": "Cable", "category": "Parts", "quantity": 3, "location": "B-2"}, headers={"X-CSRF-Token": token})
    added = client.post("/api/products/ABC123/quantity", json={"action": "ADD", "amount": 2}, headers={"X-CSRF-Token": token})
    assert added.json["product"]["quantity"] == 5
    removed = client.post("/api/products/ABC123/quantity", json={"action": "REMOVE", "amount": 4}, headers={"X-CSRF-Token": token})
    assert removed.json["product"]["quantity"] == 1
    rejected = client.post("/api/products/ABC123/quantity", json={"action": "REMOVE", "amount": 2}, headers={"X-CSRF-Token": token})
    assert rejected.status_code == 400
    assert app.extensions["excel_store"].find_product("ABC123")["quantity"] == 1


def test_api_authorization_and_csrf(client):
    assert client.get("/api/products/X").status_code == 401
    login(client)
    assert client.post("/api/products", json={}).status_code == 400


def test_workbook_has_required_sheets(app):
    wb = load_workbook(app.config["WORKBOOK_PATH"])
    assert wb.sheetnames == ["Products", "Transactions", "Users"]

