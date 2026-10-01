import os
import re
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from filelock import FileLock
from openpyxl import Workbook, load_workbook
from werkzeug.security import check_password_hash, generate_password_hash


PRODUCT_HEADERS = ["Product ID", "Barcode", "Product Name", "Category", "Description", "Quantity", "Minimum Quantity", "Location", "Date Created", "Last Updated"]
TRANSACTION_HEADERS = ["Transaction ID", "Product ID", "Product Name", "Action", "Quantity Changed", "Previous Quantity", "New Quantity", "User", "Timestamp", "Barcode"]
USER_HEADERS = ["User ID", "Full Name", "Username", "Password Hash", "Role", "Active", "Date Created", "Last Login", "Must Change Password", "Failed Attempts", "Locked Until"]
LOGIN_HEADERS = ["User", "Timestamp", "Successful", "Event", "IP Address"]


class StoreError(Exception): pass
class NotFoundError(StoreError): pass
class ValidationError(StoreError): pass
class ConflictError(StoreError): pass


class ExcelStore:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = FileLock(str(self.path) + ".lock", timeout=15)
        self._initialize()

    @staticmethod
    def _now(): return datetime.now(timezone.utc).isoformat(timespec="seconds")

    def _initialize(self):
        with self.lock:
            if not self.path.exists():
                wb = Workbook(); wb.active.title = "Products"
                wb.create_sheet("Transactions"); wb.create_sheet("Users"); wb.create_sheet("Login History")
            else:
                wb = load_workbook(self.path)
            definitions = {"Products": PRODUCT_HEADERS, "Transactions": TRANSACTION_HEADERS, "Users": USER_HEADERS, "Login History": LOGIN_HEADERS}
            for name, headers in definitions.items():
                ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
                existing = [c.value for c in ws[1]] if ws.max_row else []
                if not any(existing):
                    if ws.max_row: ws.delete_rows(1, ws.max_row)
                    ws.append(headers)
                else:
                    # Preserve old records and append new fields when upgrading an existing workbook.
                    for header in headers:
                        if header not in existing:
                            ws.cell(1, ws.max_column + 1, header)
                            existing.append(header)
                    if name == "Products":
                        aliases = {"Name": "Product Name", "Created At": "Date Created", "Updated At": "Last Updated"}
                        for old, new in aliases.items():
                            if old in existing and new in existing:
                                for row in range(2, ws.max_row + 1):
                                    if ws.cell(row, existing.index(new)+1).value is None: ws.cell(row, existing.index(new)+1, ws.cell(row, existing.index(old)+1).value)
                    if name == "Users":
                        aliases = {"Username": "Username", "Password Hash": "Password Hash", "Role": "Role", "Enabled": "Active", "Created At": "Date Created"}
                        for old, new in aliases.items():
                            if old in existing and new in existing and old != new:
                                for row in range(2, ws.max_row + 1):
                                    if ws.cell(row, existing.index(new)+1).value is None: ws.cell(row, existing.index(new)+1, ws.cell(row, existing.index(old)+1).value)
            # Backfill migration defaults and IDs.
            pws, uws = wb["Products"], wb["Users"]
            for row in range(2, pws.max_row + 1):
                if self._cell_by_header(pws,row,"Minimum Quantity") is None: self._set_by_header(pws,row,"Minimum Quantity",5)
                if self._cell_by_header(pws,row,"Description") is None: self._set_by_header(pws,row,"Description","")
            for row in range(2, uws.max_row + 1):
                username = self._cell_by_header(uws, row, "Username")
                if username and not self._cell_by_header(uws, row, "User ID"):
                    self._set_by_header(uws, row, "User ID", f"USR-{row-1:06d}")
                if username and not self._cell_by_header(uws, row, "Full Name"):
                    self._set_by_header(uws, row, "Full Name", str(username).replace(".", " ").title())
                if self._cell_by_header(uws, row, "Role") == "User": self._set_by_header(uws, row, "Role", "Staff")
                for field, value in (("Must Change Password", False), ("Failed Attempts", 0)):
                    if self._cell_by_header(uws, row, field) is None: self._set_by_header(uws, row, field, value)
            self._save_atomic(wb)

    @staticmethod
    def _cell_by_header(ws, row, header):
        headers = [c.value for c in ws[1]]
        return ws.cell(row, headers.index(header)+1).value if header in headers else None

    @staticmethod
    def _set_by_header(ws, row, header, value):
        headers = [c.value for c in ws[1]]
        ws.cell(row, headers.index(header)+1, value)

    @staticmethod
    def _rows(ws):
        headers = [cell.value for cell in ws[1]]
        for values in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in values): yield dict(zip(headers, values))

    @staticmethod
    def _append_record(ws, record):
        headers=[c.value for c in ws[1]]
        ws.append([record.get(header) for header in headers])

    def _load(self): return load_workbook(self.path)

    def _save_atomic(self, wb):
        fd, temp_name = tempfile.mkstemp(suffix=".xlsx", dir=self.path.parent); os.close(fd)
        try: wb.save(temp_name); os.replace(temp_name, self.path)
        finally:
            if os.path.exists(temp_name): os.unlink(temp_name)

    @staticmethod
    def _clean_barcode(value):
        barcode = str(value or "").strip()
        if not barcode or len(barcode) > 128: raise ValidationError("Barcode is required and must be 128 characters or fewer.")
        return barcode

    @staticmethod
    def _public_product(row):
        return {"product_id": row.get("Product ID"), "barcode": str(row.get("Barcode", "")), "name": row.get("Product Name", row.get("Name", "")), "category": row.get("Category", ""), "description": row.get("Description", "") or "", "quantity": int(row.get("Quantity") or 0), "minimum_quantity": int(row.get("Minimum Quantity") or 5), "location": row.get("Location", ""), "created_at": row.get("Date Created", row.get("Created At")), "updated_at": row.get("Last Updated", row.get("Updated At"))}

    @staticmethod
    def _public_user(row):
        return {"user_id": row.get("User ID"), "name": row.get("Full Name") or row.get("Username"), "username": row.get("Username"), "role": "Staff" if row.get("Role") == "User" else row.get("Role"), "enabled": bool(row.get("Active", row.get("Enabled", True))), "created_at": row.get("Date Created", row.get("Created At")), "last_login": row.get("Last Login"), "must_change_password": bool(row.get("Must Change Password"))}

    def authenticate(self, username, password, ip=""):
        username = str(username or "").strip().lower(); now = datetime.now(timezone.utc)
        with self.lock:
            wb = self._load(); uws = wb["Users"]; match = None; row_num = None
            for n in range(2, uws.max_row+1):
                if str(self._cell_by_header(uws,n,"Username") or "").lower() == username: match, row_num = dict(zip([c.value for c in uws[1]], [c.value for c in uws[n]])), n; break
            valid = False
            if match:
                lock_until = match.get("Locked Until")
                locked = False
                if lock_until:
                    try: locked = datetime.fromisoformat(str(lock_until)).replace(tzinfo=timezone.utc) > now
                    except ValueError: pass
                try: valid = bool(match.get("Active", match.get("Enabled", True))) and not locked and check_password_hash(match.get("Password Hash", ""), password or "")
                except (ValueError, TypeError): valid = False
            self._append_login(wb, match.get("Username") if match else (username or "Unknown"), bool(valid), "LOGIN", ip)
            if match:
                if valid:
                    self._set_by_header(uws,row_num,"Failed Attempts",0); self._set_by_header(uws,row_num,"Locked Until",None); self._set_by_header(uws,row_num,"Last Login",self._now())
                else:
                    attempts = int(match.get("Failed Attempts") or 0) + 1; self._set_by_header(uws,row_num,"Failed Attempts",attempts)
                    if attempts >= 5: self._set_by_header(uws,row_num,"Locked Until",(now+timedelta(minutes=15)).isoformat(timespec="seconds"))
            self._save_atomic(wb)
            return self._public_user(match) if valid else None

    def _append_login(self, wb, username, success, event, ip):
        self._append_record(wb["Login History"], {"User":username,"Timestamp":self._now(),"Successful":success,"Event":event,"IP Address":str(ip or "")[:64]})

    def record_logout(self, username, ip=""):
        with self.lock:
            wb=self._load(); self._append_login(wb,username,True,"LOGOUT",ip); self._save_atomic(wb)

    def record_reset_request(self, requester, ip=""):
        with self.lock:
            wb=self._load(); self._append_login(wb,str(requester or "Not provided")[:100],False,"PASSWORD RESET REQUEST",ip); self._save_atomic(wb)

    def get_user(self, username):
        with self.lock:
            wb=self._load()
            for r in self._rows(wb["Users"]):
                if str(r.get("Username","")).lower()==str(username).lower(): return self._public_user(r)
        return None

    def list_users(self):
        with self.lock: return [self._public_user(r) for r in self._rows(self._load()["Users"])]

    def list_login_activity(self, limit=100):
        with self.lock:
            rows=list(self._rows(self._load()["Login History"]))[-limit:]; rows.reverse(); return rows

    def create_user(self, username, password, role="Staff", name=None, temporary=False):
        username=str(username or "").strip().lower(); role=str(role or "Staff").title(); name=str(name or username).strip()
        if role == "User": role = "Staff"
        if not re.fullmatch(r"[a-z0-9_.-]{3,40}",username): raise ValidationError("Username must be 3–40 characters using letters, numbers, dots, dashes, or underscores.")
        if len(password or "") < 10: raise ValidationError("Password must be at least 10 characters.")
        if role not in {"Admin","Staff"}: raise ValidationError("Role must be Admin or Staff.")
        if not name or len(name)>100: raise ValidationError("Full name is required and must be 100 characters or fewer.")
        with self.lock:
            wb=self._load(); ws=wb["Users"]
            if any(str(r.get("Username","")).lower()==username for r in self._rows(ws)): raise ConflictError("That username already exists.")
            uid=max([int(str(r.get("User ID", "USR-000000")).split("-")[-1]) for r in self._rows(ws) if str(r.get("User ID", "")).startswith("USR-")] or [0])+1
            self._append_record(ws,{"User ID":f"USR-{uid:06d}","Full Name":name,"Username":username,"Password Hash":generate_password_hash(password),"Role":role,"Active":True,"Date Created":self._now(),"Last Login":None,"Must Change Password":bool(temporary),"Failed Attempts":0,"Locked Until":None}); self._save_atomic(wb)
        return self.get_user(username)

    def update_user(self, username, role=None, enabled=None, password=None):
        with self.lock:
            wb=self._load(); ws=wb["Users"]
            for n in range(2,ws.max_row+1):
                if str(self._cell_by_header(ws,n,"Username")).lower()==str(username).lower():
                    if role is not None:
                        role=str(role).title(); role="Staff" if role=="User" else role
                        if role not in {"Admin","Staff"}: raise ValidationError("Role must be Admin or Staff.")
                        self._set_by_header(ws,n,"Role",role)
                    if enabled is not None: self._set_by_header(ws,n,"Active",bool(enabled))
                    if password:
                        if len(password)<10: raise ValidationError("Password must be at least 10 characters.")
                        self._set_by_header(ws,n,"Password Hash",generate_password_hash(password)); self._set_by_header(ws,n,"Must Change Password",True); self._set_by_header(ws,n,"Failed Attempts",0); self._set_by_header(ws,n,"Locked Until",None)
                    self._save_atomic(wb); return
            raise NotFoundError("User not found.")

    def change_password(self, username, current, new):
        if len(new or "") < 10: raise ValidationError("New password must be at least 10 characters.")
        with self.lock:
            wb=self._load(); ws=wb["Users"]
            for n in range(2,ws.max_row+1):
                if str(self._cell_by_header(ws,n,"Username")).lower()==str(username).lower():
                    hashed=self._cell_by_header(ws,n,"Password Hash")
                    if not check_password_hash(hashed,current or ""): raise ValidationError("Current password is incorrect.")
                    self._set_by_header(ws,n,"Password Hash",generate_password_hash(new)); self._set_by_header(ws,n,"Must Change Password",False); self._save_atomic(wb); return
            raise NotFoundError("User not found.")

    def list_products(self):
        with self.lock: return [self._public_product(r) for r in self._rows(self._load()["Products"])]

    def find_product(self, barcode):
        barcode=self._clean_barcode(barcode)
        with self.lock:
            for r in self._rows(self._load()["Products"]):
                if str(r.get("Barcode"))==barcode: return self._public_product(r)
        return None

    def find_product_query(self, query):
        q=str(query or "").strip().lower()
        if not q: return []
        return [p for p in self.list_products() if q in " ".join(str(p.get(k) or "") for k in ("product_id","barcode","name","category","description","location")).lower()]

    def create_product(self, data, username):
        barcode=self._clean_barcode(data.get("barcode")); name=str(data.get("name") or data.get("product_name") or "").strip(); category=str(data.get("category") or "").strip(); location=str(data.get("location") or "").strip(); description=str(data.get("description") or "").strip()
        try: quantity=int(data.get("quantity",0)); minimum=int(data.get("minimum_quantity",data.get("minimum",5)))
        except (ValueError,TypeError): raise ValidationError("Quantity values must be whole numbers.")
        if not name or len(name)>120 or not category or len(category)>80 or not location or len(location)>80 or len(description)>500: raise ValidationError("Name, category, and location are required. Check the field limits.")
        if quantity<0 or minimum<0: raise ValidationError("Quantities cannot be negative.")
        with self.lock:
            wb=self._load(); ws=wb["Products"]; rows=list(self._rows(ws))
            if any(str(r.get("Barcode"))==barcode for r in rows): raise ConflictError("A product with that barcode already exists.")
            max_id=max([int(str(r.get("Product ID")).split("-")[-1]) for r in rows if re.fullmatch(r"PRD-\d+",str(r.get("Product ID")))] or [0]); pid=f"PRD-{max_id+1:06d}"; now=self._now()
            self._append_record(ws,{"Product ID":pid,"Barcode":barcode,"Product Name":name,"Name":name,"Category":category,"Description":description,"Quantity":quantity,"Minimum Quantity":minimum,"Location":location,"Date Created":now,"Created At":now,"Last Updated":now,"Updated At":now}); self._append_transaction(wb,pid,name,"CREATE",quantity,0,quantity,username,barcode,now); self._save_atomic(wb)
        return self.find_product(barcode)

    def update_product(self, barcode, data, username):
        with self.lock:
            wb=self._load(); ws=wb["Products"]
            for n in range(2,ws.max_row+1):
                if str(ws.cell(n,2).value)==str(barcode):
                    fields={"Product Name":str(data.get("name","")).strip(),"Category":str(data.get("category","")).strip(),"Description":str(data.get("description","")).strip(),"Minimum Quantity":data.get("minimum_quantity",5),"Location":str(data.get("location","")).strip()}
                    try: fields["Minimum Quantity"]=int(fields["Minimum Quantity"])
                    except (ValueError,TypeError): raise ValidationError("Minimum quantity must be a whole number.")
                    if not fields["Product Name"] or not fields["Category"] or not fields["Location"] or fields["Minimum Quantity"]<0: raise ValidationError("Enter valid product details and a non-negative minimum quantity.")
                    for field,value in fields.items(): self._set_by_header(ws,n,field,value)
                    now=self._now(); self._set_by_header(ws,n,"Last Updated",now)
                    quantity=int(self._cell_by_header(ws,n,"Quantity") or 0)
                    self._append_transaction(wb,self._cell_by_header(ws,n,"Product ID"),fields["Product Name"],"EDIT",0,quantity,quantity,username,barcode,now)
                    self._save_atomic(wb); return self.find_product(barcode)
        raise NotFoundError("Product not found.")

    def _append_transaction(self, wb, pid, name, action, amount, previous, new, username, barcode, now=None):
        ws=wb["Transactions"]; n=ws.max_row; stamp=now or self._now(); self._append_record(ws,{"Transaction ID":f"TXN-{n:06d}","Product ID":pid,"Product Name":name,"Action":action,"Quantity Changed":amount,"Previous Quantity":previous,"New Quantity":new,"User":username,"Timestamp":stamp,"Barcode":barcode,"Date":stamp,"Amount":amount,"Updated Quantity":new})

    def update_quantity(self, barcode, action, amount, username):
        barcode=self._clean_barcode(barcode); action=str(action or "").upper()
        try: amount=int(amount)
        except (ValueError,TypeError): raise ValidationError("Amount must be a whole number.")
        if action not in {"ADD","REMOVE"} or amount<=0: raise ValidationError("Choose add or remove and enter an amount greater than zero.")
        with self.lock:
            wb=self._load(); ws=wb["Products"]
            for n in range(2,ws.max_row+1):
                if str(ws.cell(n,2).value)==barcode:
                    old=int(self._cell_by_header(ws,n,"Quantity") or 0); new=old+amount if action=="ADD" else old-amount
                    if new<0: raise ValidationError("Cannot remove more stock than is available.")
                    now=self._now(); self._set_by_header(ws,n,"Quantity",new); self._set_by_header(ws,n,"Last Updated",now); pid=self._cell_by_header(ws,n,"Product ID"); name=self._cell_by_header(ws,n,"Product Name") or self._cell_by_header(ws,n,"Name")
                    self._append_transaction(wb,pid,name,action,amount,old,new,username,barcode,now); self._save_atomic(wb); return self.find_product(barcode)
        raise NotFoundError("Product not found.")

    def list_transactions(self, limit=100, filters=None):
        filters=filters or {}
        with self.lock: rows=list(self._rows(self._load()["Transactions"]))
        result=[]
        for r in rows:
            r.setdefault("Date",r.get("Timestamp")); r.setdefault("User",r.get("User")); r.setdefault("Action",""); r.setdefault("Amount",r.get("Quantity Changed",0)); r.setdefault("Updated Quantity",r.get("New Quantity",0)); r.setdefault("Product ID",r.get("Product ID"))
            pairs=(("product",("Product Name","Product ID","Barcode")),("user",("User",)),("action",("Action",)),("date",("Timestamp","Date")))
            if all(not filters.get(key) or any(str(r.get(col) or "").lower().find(str(filters[key]).lower())>=0 for col in cols) for key,cols in pairs): result.append(r)
        result.reverse(); return result[:limit] if limit else result
